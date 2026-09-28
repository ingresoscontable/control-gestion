from io import BytesIO
from urllib.parse import unquote

from helpers import (
    _id_meta,
    login,
)
from openpyxl import load_workbook

from app import backup, config
from app.database import SessionLocal
from app.models import Auditoria, Meta


def test_reporte_pdf_incluye_el_avance(client, monkeypatch):
    import fpdf

    original = fpdf.FPDF.__init__

    def sin_compresion(self, *args, **kwargs):
        original(self, *args, **kwargs)
        self.set_compression(False)  # deja el texto legible para poder verificarlo

    monkeypatch.setattr(fpdf.FPDF, "__init__", sin_compresion)

    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={"titulo": "Meta del PDF", "horas_estimadas": "8"},
        follow_redirects=False,
    )

    respuesta = client.get("/reportes/mensual.pdf")
    assert respuesta.status_code == 200

    texto = respuesta.content.decode("latin-1")
    assert "Avance promedio de metas activas" in texto
    assert "Metas con plazo vencido" in texto
    assert "Meta del PDF" in texto
    # Encabezados de la tabla de metas, incluida la columna nueva de avance.
    # La columna dice "periodo" y no "mes" porque el mismo PDF sirve tambien
    # para los informes por rango de fechas libre.
    for encabezado in ("Asignada a", "Plazo", "Avances periodo", "Estado"):
        assert encabezado in texto


def test_exportar_excel_con_filtros(client):
    assert login(client, "jefe", "1234").status_code == 303

    respuesta = client.get("/registros/exportar.xlsx")
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in respuesta.headers["content-disposition"]
    assert respuesta.content[:2] == b"PK"  # un .xlsx es un zip

    hoja = load_workbook(BytesIO(respuesta.content)).active
    assert hoja["A1"].value == "Fecha"
    textos = [c.value for fila in hoja.iter_rows() for c in fila if c.value]
    assert "Revisadas 40 facturas pendientes" in textos
    assert "TOTAL HORAS" in textos

    # Filtrando por una persona sin registros, el archivo sale sin datos.
    filtrado = client.get("/registros/exportar.xlsx", params={"usuario_id": "999"})
    valores = [
        c.value
        for fila in load_workbook(BytesIO(filtrado.content)).active.iter_rows()
        for c in fila
        if c.value
    ]
    assert "Revisadas 40 facturas pendientes" not in valores

    # Un empleado no puede exportar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert client.get(
        "/registros/exportar.xlsx", follow_redirects=False
    ).headers["location"].startswith("/?msg=")


def test_redireccion_sin_json_en_el_cuerpo(client):
    """La redireccion de seguridad responde 303 con Location y sin cuerpo."""
    respuesta = client.get("/", follow_redirects=False)
    assert respuesta.status_code == 303
    assert respuesta.headers["location"] == "/login"
    assert "application/json" not in respuesta.headers.get("content-type", "")
    assert respuesta.content == b""


def test_clave_de_sesion_propia_por_instalacion(tmp_path, monkeypatch):
    """Sin CG_SECRET_KEY se genera una clave por instalacion y se reutiliza."""
    monkeypatch.delenv("CG_SECRET_KEY", raising=False)
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)

    primera = config._clave_de_sesion()
    assert len(primera) >= 32
    assert config._clave_de_sesion() == primera
    assert (tmp_path / ".secret_key").read_text(encoding="utf-8") == primera

    # Si el ambiente define la clave, manda la del ambiente.
    monkeypatch.setenv("CG_SECRET_KEY", "la-que-diga-el-ambiente")
    assert config._clave_de_sesion() == "la-que-diga-el-ambiente"


def test_aviso_si_el_respaldo_queda_en_el_mismo_disco(client):
    # Por defecto los respaldos van a data/respaldos: mismo disco que la base.
    assert backup.respaldos_en_mismo_disco() is True
    assert login(client, "jefe", "1234").status_code == 303
    assert "mismo disco" in client.get("/respaldos").text


def test_restaurar_respaldo(client):
    assert login(client, "jefe", "1234").status_code == 303

    client.post(
        "/metas", data={"titulo": "Meta que va a volver"}, follow_redirects=False
    )
    client.post("/respaldos/ahora", follow_redirects=False)
    respaldo = max(
        config.BACKUP_DIR.glob("control_*_manual.db"),
        key=lambda ruta: ruta.stat().st_mtime,
    )

    meta_id = _id_meta("Meta que va a volver")
    client.post(f"/metas/{meta_id}/eliminar", follow_redirects=False)

    db = SessionLocal()
    try:
        archivada = db.query(Meta).filter_by(titulo="Meta que va a volver").one()
        assert archivada.estado == "eliminada"
    finally:
        db.close()

    respuesta = client.post(
        "/respaldos/restaurar", data={"nombre": respaldo.name}, follow_redirects=False
    )
    assert respuesta.status_code == 303
    assert "Base restaurada" in unquote(respuesta.headers["location"])

    db = SessionLocal()
    try:
        assert db.query(Meta).filter_by(titulo="Meta que va a volver").count() == 1
        # La fila de auditoria se escribe despues de restaurar: si fuera antes,
        # la restauracion la tiraria junto con el resto de la base vieja.
        assert db.query(Auditoria).filter_by(accion="restaurar_respaldo").count() == 1
    finally:
        db.close()

    # La pagina lista el respaldo y la copia de seguridad previa.
    assert "antes-de-restaurar" in client.get("/respaldos").text

    # Un empleado no puede restaurar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    denegado = client.post(
        "/respaldos/restaurar", data={"nombre": respaldo.name}, follow_redirects=False
    )
    assert denegado.status_code == 303
    assert denegado.headers["location"].startswith("/?msg=")


def test_restaurar_rechaza_archivos_que_no_son_respaldos(client):
    assert login(client, "jefe", "1234").status_code == 303

    (config.BACKUP_DIR / "pepito.txt").write_text("hola", encoding="utf-8")
    (config.BACKUP_DIR / "control_basura.db").write_text(
        "esto no es sqlite", encoding="utf-8"
    )

    casos = [
        "../../control.db",  # fuera de la carpeta de respaldos
        "no_existe.db",  # no esta
        "pepito.txt",  # esta pero no tiene formato de respaldo
        "control_basura.db",  # tiene el formato pero esta corrupto
    ]
    for nombre in casos:
        respuesta = client.post(
            "/respaldos/restaurar", data={"nombre": nombre}, follow_redirects=False
        )
        assert respuesta.status_code == 303
        assert "No se pudo restaurar" in unquote(respuesta.headers["location"]), nombre


