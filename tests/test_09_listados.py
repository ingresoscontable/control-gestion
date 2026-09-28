import calendar
from datetime import date, timedelta
from io import BytesIO

from helpers import (
    _id_meta,
    _ultimo_registro_de,
    login,
)
from openpyxl import load_workbook

from app.database import SessionLocal
from app.metricas import calendario_mes
from app.models import Meta, Registro, Usuario


def test_filtrar_registros_por_meta(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas", data={"titulo": "Meta del filtro A", "horas_estimadas": "10"}
    )
    client.post(
        "/metas", data={"titulo": "Meta del filtro B", "horas_estimadas": "10"}
    )
    meta_a = _id_meta("Meta del filtro A")
    meta_b = _id_meta("Meta del filtro B")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    for meta_id, texto in ((meta_a, "Carga de la meta A"), (meta_b, "Carga de la meta B")):
        assert (
            client.post(
                "/registros",
                data={"meta_id": str(meta_id), "descripcion": texto, "horas": "2"},
                follow_redirects=False,
            ).status_code
            == 303
        )

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303

    todo = client.get("/registros").text
    assert "Carga de la meta A" in todo
    assert "Carga de la meta B" in todo

    # Filtrando por meta A solo queda esa carga y el combo marca la seleccion.
    pagina = client.get("/registros", params={"meta_id": str(meta_a)})
    assert pagina.status_code == 200
    assert "Carga de la meta A" in pagina.text
    assert "Carga de la meta B" not in pagina.text
    assert f'value="{meta_a}" selected' in pagina.text

    # La exportacion respeta el mismo filtro.
    hoja = load_workbook(
        BytesIO(
            client.get(
                "/registros/exportar.xlsx", params={"meta_id": str(meta_a)}
            ).content
        )
    ).active
    textos = [c.value for fila in hoja.iter_rows() for c in fila if c.value]
    assert "Carga de la meta A" in textos
    assert "Carga de la meta B" not in textos

    # Un meta_id que no es numero se ignora en vez de romper.
    assert client.get("/registros", params={"meta_id": "abc"}).status_code == 200


def test_duplicar_meta(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={
            "titulo": "Corte de septiembre 2026",
            "descripcion": "Cierra las cuentas del mes",
            "fecha_inicio": "2026-09-01",
            "fecha_limite": "2026-09-30",
            "horas_estimadas": "12",
        },
    )
    original = _id_meta("Corte de septiembre 2026")

    # El original se cierra: la copia tiene que arrancar activa igual.
    client.post(f"/metas/{original}/estado", data={"estado": "cerrada"})

    respuesta = client.post(f"/metas/{original}/duplicar", follow_redirects=False)
    assert respuesta.status_code == 303
    destino = respuesta.headers["location"]
    assert destino.startswith("/metas/") and "/editar?msg=" in destino
    assert "Meta%20duplicada" in destino

    db = SessionLocal()
    try:
        copia = db.get(Meta, int(destino.split("/")[2]))
        assert copia.id != original
        assert copia.titulo == "Corte de septiembre 2026 (copia)"
        assert copia.descripcion == "Cierra las cuentas del mes"
        assert copia.fecha_inicio == date(2026, 9, 1)
        assert copia.fecha_limite == date(2026, 9, 30)
        assert copia.horas_estimadas == 12
        assert copia.estado == "activa"
        assert db.get(Meta, original).estado == "cerrada"
    finally:
        db.close()

    assert "Corte de septiembre 2026 (copia)" in client.get("/metas").text

    # Un empleado no puede duplicar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    denegado = client.post(f"/metas/{original}/duplicar", follow_redirects=False)
    assert denegado.status_code == 303
    assert denegado.headers["location"].startswith("/?msg=")

    # Archivar no borra: la meta queda en pie y de paso se puede duplicar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    borrada = client.post(f"/metas/{original}/eliminar", follow_redirects=False)
    assert borrada.status_code == 303

    db = SessionLocal()
    try:
        assert db.get(Meta, original).estado == "eliminada"
    finally:
        db.close()

    copia_archivada = client.post(f"/metas/{original}/duplicar", follow_redirects=False)
    assert copia_archivada.status_code == 303
    assert "Meta%20duplicada" in copia_archivada.headers["location"]

    # Duplicar algo que ya no existe no tira error.
    inexistente = client.post("/metas/999999/duplicar", follow_redirects=False)
    assert inexistente.status_code == 303
    assert "Esa%20meta%20ya%20no%20existe" in inexistente.headers["location"]


def test_repetir_carga_del_dia_anterior(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post("/metas", data={"titulo": "Meta para repetir", "horas_estimadas": "8"})
    meta_id = _id_meta("Meta para repetir")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    client.post(
        "/registros",
        data={
            "meta_id": str(meta_id),
            "descripcion": "Trabajo de ayer para repetir",
            "fecha": (date.today() - timedelta(days=1)).isoformat(),
            "horas": "6.5",
            "estado": "completado",
        },
    )
    registro_id = _ultimo_registro_de("mperez")

    # El formulario queda precargado con lo anterior pero la fecha es hoy.
    panel = client.get("/", params={"repetir": str(registro_id)})
    assert panel.status_code == 200
    assert "Repitiendo el registro del" in panel.text
    assert "required>Trabajo de ayer para repetir</textarea>" in panel.text
    assert 'value="6.5"' in panel.text
    assert f'value="{meta_id}" selected' in panel.text
    assert f'value="{date.today().isoformat()}"' in panel.text
    assert f'href="/?repetir={registro_id}"' in panel.text

    # Se guarda como un registro nuevo del dia de hoy.
    assert (
        client.post(
            "/registros",
            data={
                "meta_id": str(meta_id),
                "descripcion": "Trabajo de ayer para repetir",
                "horas": "6.5",
                "estado": "completado",
            },
            follow_redirects=False,
        ).headers["location"].startswith("/?msg=Registro%20guardado")
    )
    db = SessionLocal()
    try:
        persona = db.query(Usuario).filter_by(usuario="mperez").one()
        hoyos = (
            db.query(Registro)
            .filter(
                Registro.usuario_id == persona.id,
                Registro.fecha == date.today(),
                Registro.meta_id == meta_id,
            )
            .count()
        )
        totales = (
            db.query(Registro)
            .filter(
                Registro.usuario_id == persona.id,
                Registro.meta_id == meta_id,
            )
            .count()
        )
        # Solo la copia es de hoy: la original quedo en el dia anterior.
        assert hoyos == 1
        assert totales == 2
    finally:
        db.close()

    # Un id inexistente o ajeno no precarga nada ni rompe la pagina.
    assert "Repitiendo el registro del" not in client.get(
        "/", params={"repetir": "999999"}
    ).text
    assert "Repitiendo el registro del" not in client.get(
        "/", params={"repetir": "no-numero"}
    ).text

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    ajeno = client.get("/", params={"repetir": str(registro_id)})
    assert ajeno.status_code == 200
    assert "Repitiendo el registro del" not in ajeno.text
    # El jefe solo ve el boton de repetir en sus propios registros.
    assert f'href="/?repetir={registro_id}"' not in ajeno.text


def test_calendario_de_carga(client):
    hoy = date.today()

    db = SessionLocal()
    try:
        equipo = (
            db.query(Usuario).filter_by(activo=True).order_by(Usuario.nombre).all()
        )
        cal = calendario_mes(db, equipo, hoy.year, hoy.month)
        assert cal["huecos"] == date(hoy.year, hoy.month, 1).weekday()
        assert len(cal["dias"]) == calendar.monthrange(hoy.year, hoy.month)[1]
        assert cal["total_horas"] >= 0
        del_hoy = next(d for d in cal["dias"] if d["fecha"] == hoy)
        assert len(del_hoy["celdas"]) == len(equipo)
        assert all("iniciales" in celda for celda in del_hoy["celdas"])
    finally:
        db.close()

    # Un empleado no entra.
    assert login(client, "mperez", "4821").status_code == 303
    denegado = client.get("/calendario", follow_redirects=False)
    assert denegado.status_code == 303
    assert denegado.headers["location"].startswith("/?msg=")
    client.post("/logout", follow_redirects=False)

    assert login(client, "jefe", "1234").status_code == 303
    pagina = client.get("/calendario", params={"anio": hoy.year, "mes": hoy.month})
    assert pagina.status_code == 200
    assert "Calendario de carga" in pagina.text
    assert ">Lun<" in pagina.text
    assert ">Dom<" in pagina.text
    assert "Mes anterior" in pagina.text
    assert "Mes siguiente" in pagina.text
    assert "cal-body" in pagina.text

    # Se puede ir a cualquier mes.
    agosto = client.get("/calendario", params={"anio": 2026, "mes": 8})
    assert agosto.status_code == 200
    assert '<option value="8" selected>agosto</option>' in agosto.text
    db = SessionLocal()
    try:
        equipo = (
            db.query(Usuario).filter_by(activo=True).order_by(Usuario.nombre).all()
        )
        agosto_cal = calendario_mes(db, equipo, 2026, 8)
    finally:
        db.close()
    # Celdas dibujadas = huecos del lunes + dias del mes.
    assert agosto.text.count('class="cal-dia') == (
        agosto_cal["huecos"] + 31
    )

    # Periodo raro: si no es numero vuelve al mes en curso, si esta fuera de
    # rango avisa.
    assert client.get(
        "/calendario", params={"anio": "abc"}, follow_redirects=False
    ).status_code == 200
    fuera = client.get(
        "/calendario", params={"mes": "13"}, follow_redirects=False
    )
    assert fuera.status_code == 303
    assert fuera.headers["location"].startswith("/calendario?msg=")


