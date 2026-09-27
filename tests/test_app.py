"""Pruebas de la aplicación de Control de Gestión."""

import calendar
import os
import tempfile
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

# Se define la carpeta de datos antes de importar la app para no tocar la
# base real de la oficina.
os.environ["CG_DATA_DIR"] = tempfile.mkdtemp(prefix="cg-test-")
os.environ["CG_SECRET_KEY"] = "clave-de-prueba"
os.environ["CG_ADMIN_USUARIO"] = "jefe"
os.environ["CG_ADMIN_PIN"] = "1234"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import backup, config, migraciones, seed  # noqa: E402
from app.database import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.metricas import calendario_mes  # noqa: E402
from app.models import Meta, Registro, Usuario  # noqa: E402
from app.security import hash_pin  # noqa: E402


@pytest.fixture()
def client():
    with TestClient(app) as c:
        yield c


def login(client: TestClient, usuario: str, pin: str):
    return client.post(
        "/login", data={"usuario": usuario, "pin": pin}, follow_redirects=False
    )


def test_login_sin_sesion_redirige(client):
    respuesta = client.get("/", follow_redirects=False)
    assert respuesta.status_code == 303
    assert respuesta.headers["location"] == "/login"


def test_pin_incorrecto_no_inicia_sesion(client):
    respuesta = login(client, "jefe", "0000")
    assert respuesta.status_code == 401
    assert "incorrecto" in respuesta.text


def test_pin_se_guarda_hasheado(client):
    db = SessionLocal()
    try:
        jefe = db.query(Usuario).filter_by(usuario="jefe").one()
        assert jefe.pin_hash != "1234"
        assert jefe.pin_hash.startswith("pbkdf2$")
    finally:
        db.close()


def test_jefe_entra_y_ve_su_panel(client):
    assert login(client, "jefe", "1234").status_code == 303
    panel = client.get("/")
    assert panel.status_code == 200
    assert "Resumen de hoy" in panel.text


def test_primer_ingreso_se_borra_cuando_el_jefe_cambia_su_pin(client):
    ruta = seed.guardar_primer_ingreso("jefe", "9876")
    assert ruta.is_file()
    assert "9876" in ruta.read_text(encoding="utf-8")

    login(client, "jefe", "1234")
    cambio = client.post(
        "/mi-pin",
        data={"pin_actual": "1234", "pin_nuevo": "5678"},
        follow_redirects=False,
    )
    assert cambio.status_code == 303
    assert not ruta.exists()

    db = SessionLocal()
    try:
        jefe = db.query(Usuario).filter_by(usuario="jefe").one()
        jefe.pin_hash = hash_pin("1234")
        db.commit()
    finally:
        db.close()


def test_flujo_completo_jefe_y_empleado(client):
    assert login(client, "jefe", "1234").status_code == 303

    # El jefe crea dos accesos, uno con PIN inválido (debe rebotar).
    assert (
        client.post(
            "/equipo",
            data={"nombre": "Maria Perez", "usuario": "mperez", "pin": "12", "rol": "empleado"},
            follow_redirects=False,
        ).headers["location"]
        .startswith("/equipo")
    )
    assert client.post(
        "/equipo",
        data={
            "nombre": "Maria Perez",
            "usuario": "mperez",
            "cargo": "Analista contable",
            "pin": "4821",
            "rol": "empleado",
        },
        follow_redirects=False,
    ).status_code == 303

    # El jefe crea una meta asignada a Maria.
    db = SessionLocal()
    try:
        maria = db.query(Usuario).filter_by(usuario="mperez").one()
        maria_id = maria.id
    finally:
        db.close()

    assert client.post(
        "/metas",
        data={
            "titulo": "Conciliacion bancaria septiembre",
            "descripcion": "Cerrar la conciliacion del mes",
            "asignado_a": str(maria_id),
            "fecha_limite": "2026-10-05",
        },
        follow_redirects=False,
    ).status_code == 303
    assert "Conciliacion bancaria septiembre" in client.get("/metas").text

    # Maria entra y registra su trabajo del dia.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303

    panel = client.get("/")
    assert "Conciliacion bancaria septiembre" in panel.text

    assert client.post(
        "/registros",
        data={
            "descripcion": "Revisadas 40 facturas pendientes",
            "horas": "6,5",
            "estado": "en_progreso",
            "fecha": date.today().isoformat(),
        },
        follow_redirects=False,
    ).status_code == 303

    panel = client.get("/")
    assert "Revisadas 40 facturas pendientes" in panel.text

    # Un empleado no puede administrar metas ni ver el equipo.
    assert client.get("/metas", follow_redirects=False).headers["location"].startswith("/?msg=")
    assert client.get("/equipo", follow_redirects=False).headers["location"].startswith("/?msg=")


def test_jefe_ve_registros_y_horas(client):
    assert login(client, "jefe", "1234").status_code == 303
    respuesta = client.get("/registros")
    assert respuesta.status_code == 200
    assert "Revisadas 40 facturas pendientes" in respuesta.text
    assert "6.5 horas acumuladas" in respuesta.text


def test_texto_vacio_no_crea_registro(client):
    assert login(client, "jefe", "1234").status_code == 303
    respuesta = client.post(
        "/registros", data={"descripcion": "   "}, follow_redirects=False
    )
    assert respuesta.status_code == 303
    assert "no puede estar vacía" in unquote(respuesta.headers["location"])


def _id_usuario(login_usuario: str) -> int:
    db = SessionLocal()
    try:
        return db.query(Usuario).filter_by(usuario=login_usuario).one().id
    finally:
        db.close()


def test_alerta_de_metas_vencidas(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={
            "titulo": "Presentar declaracion atrasada",
            "asignado_a": str(_id_usuario("mperez")),
            "fecha_limite": "2020-01-31",
        },
        follow_redirects=False,
    )

    panel = client.get("/")
    assert "meta(s) con el plazo vencido" in panel.text
    assert "Presentar declaracion atrasada" in panel.text
    assert "de atraso" in panel.text

    # La analista ve el aviso, pero referido a sus metas.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert "entre las tuyas" in client.get("/").text

    # Al cerrar la meta, el aviso desaparece.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    meta_id = _id_meta("Presentar declaracion atrasada")
    client.post(
        f"/metas/{meta_id}/estado",
        data={"estado": "cerrada"},
        follow_redirects=False,
    )
    assert "Presentar declaracion atrasada" not in client.get("/").text


def test_editar_meta(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={"titulo": "Meta a editar", "horas_estimadas": "10"},
        follow_redirects=False,
    )
    meta_id = _id_meta("Meta a editar")

    formulario = client.get(f"/metas/{meta_id}/editar")
    assert formulario.status_code == 200
    assert "Meta a editar" in formulario.text
    assert "Avance actual" in formulario.text

    # Maria carga 5 horas contra las 10 previstas => 50%.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    client.post(
        "/registros",
        data={"meta_id": str(meta_id), "descripcion": "Avance", "horas": "5"},
    )

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    assert "50% · 1 reporte(s)" in client.get(f"/metas/{meta_id}/editar").text

    # Un titulo vacio no se acepta.
    assert client.post(
        f"/metas/{meta_id}", data={"titulo": "   "}, follow_redirects=False
    ).headers["location"].startswith(f"/metas/{meta_id}/editar")

    respuesta = client.post(
        f"/metas/{meta_id}",
        data={
            "titulo": "Meta editada",
            "descripcion": "Ajustada al nuevo formato",
            "asignado_a": str(_id_usuario("mperez")),
            "fecha_limite": "2030-12-31",
            "horas_estimadas": "20",
            "estado": "activa",
        },
        follow_redirects=False,
    )
    assert respuesta.status_code == 303
    assert respuesta.headers["location"] == "/metas?msg=Meta%20actualizada"

    db = SessionLocal()
    try:
        meta = db.get(Meta, meta_id)
        assert meta.titulo == "Meta editada"
        assert meta.descripcion == "Ajustada al nuevo formato"
        assert meta.horas_estimadas == 20
        assert meta.asignado_a == _id_usuario("mperez")
        assert meta.fecha_limite == date(2030, 12, 31)
    finally:
        db.close()

    # Con las mismas 5 horas pero 20 previstas, el avance baja a 25%.
    formulario = client.get(f"/metas/{meta_id}/editar")
    assert "25% · 1 reporte(s)" in formulario.text
    assert "Meta editada" in client.get("/metas").text

    # Un empleado no puede abrir el formulario de edicion.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert client.get(
        f"/metas/{meta_id}/editar", follow_redirects=False
    ).headers["location"].startswith("/?msg=")


def test_pagina_ayuda_muestra_instrucciones(client):
    assert login(client, "jefe", "1234").status_code == 303
    ayuda = client.get("/ayuda")
    assert ayuda.status_code == 200
    assert "netsh advfirewall firewall add rule" in ayuda.text


def test_resumen_semanal_en_panel_del_jefe(client):
    assert login(client, "jefe", "1234").status_code == 303
    panel = client.get("/")
    assert "Resumen de la semana" in panel.text
    assert "Maria Perez" in panel.text
    assert "de 5" in panel.text

    # El empleado no ve el resumen del equipo.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert "Resumen de la semana" not in client.get("/").text


def _id_registro_de(login_usuario: str) -> int:
    db = SessionLocal()
    try:
        persona = db.query(Usuario).filter_by(usuario=login_usuario).one()
        registro = (
            db.query(Registro)
            .filter(Registro.usuario_id == persona.id)
            .order_by(Registro.id)
            .first()
        )
        return registro.id
    finally:
        db.close()


def test_comentario_del_jefe(client):
    registro_id = _id_registro_de("mperez")

    # Un empleado no puede comentar.
    assert login(client, "mperez", "4821").status_code == 303
    assert client.post(
        f"/registros/{registro_id}/comentario",
        data={"comentario": "comentario indebido"},
        follow_redirects=False,
    ).headers["location"].startswith("/?msg=")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303

    aviso = "Buen avance, revisar el IVA del periodo."
    respuesta = client.post(
        f"/registros/{registro_id}/comentario",
        data={"comentario": aviso, "volver": "/registros"},
        follow_redirects=False,
    )
    assert respuesta.status_code == 303
    assert respuesta.headers["location"] == "/registros?msg=Comentario%20guardado"
    assert aviso in client.get("/registros").text
    assert aviso in client.get("/").text

    # El empleado lo ve en su panel (sin formulario para editar).
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    panel = client.get("/")
    assert aviso in panel.text
    assert f"/registros/{registro_id}/comentario" not in panel.text

    # Un comentario vacío lo borra.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    assert client.post(
        f"/registros/{registro_id}/comentario",
        data={"comentario": "   ", "volver": "/registros"},
        follow_redirects=False,
    ).status_code == 303
    assert aviso not in client.get("/registros").text


def test_reporte_mensual_pdf(client):
    assert login(client, "jefe", "1234").status_code == 303
    hoy = date.today()

    respuesta = client.get(
        "/reportes/mensual.pdf", params={"anio": hoy.year, "mes": hoy.month}
    )
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/pdf"
    assert respuesta.content.startswith(b"%PDF")
    assert len(respuesta.content) > 2000
    assert f"reporte_{hoy.year}-{hoy.month:02d}.pdf" in respuesta.headers[
        "content-disposition"
    ]

    # Un mes sin datos tambien genera un PDF valido.
    vacio = client.get("/reportes/mensual.pdf", params={"anio": 2000, "mes": 1})
    assert vacio.status_code == 200
    assert vacio.content.startswith(b"%PDF")

    # Periodo invalido.
    assert client.get(
        "/reportes/mensual.pdf", params={"mes": "13"}, follow_redirects=False
    ).headers["location"].startswith("/registros?msg=")

    # Un empleado no puede bajarlo.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert client.get(
        "/reportes/mensual.pdf", follow_redirects=False
    ).headers["location"].startswith("/?msg=")


def test_respaldos_automaticos_y_manuales(client):
    assert login(client, "jefe", "1234").status_code == 303

    # El respaldo automatico se crea al arrancar la aplicacion.
    automaticos = list(config.BACKUP_DIR.glob("control_*.db"))
    assert automaticos

    assert client.post("/respaldos/ahora", follow_redirects=False).status_code == 303
    todos = list(config.BACKUP_DIR.glob("control_*.db"))
    assert len(todos) == len(automaticos) + 1
    assert any(archivo.name.endswith("_manual.db") for archivo in todos)

    pagina = client.get("/respaldos")
    assert pagina.status_code == 200
    assert "Respaldos existentes" in pagina.text

    # Un empleado no entra a respaldos.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert client.get(
        "/respaldos", follow_redirects=False
    ).headers["location"].startswith("/?msg=")


def _id_meta(titulo: str) -> int:
    db = SessionLocal()
    try:
        return db.query(Meta).filter_by(titulo=titulo).one().id
    finally:
        db.close()


def test_progreso_de_metas(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={"titulo": "Cierre contable octubre", "horas_estimadas": "10"},
        follow_redirects=False,
    )
    client.post(
        "/metas", data={"titulo": "Meta sin estimacion"}, follow_redirects=False
    )
    meta_horas = _id_meta("Cierre contable octubre")
    meta_sin = _id_meta("Meta sin estimacion")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303

    # 4 horas contra una meta de 10 previstas => 40%.
    client.post(
        "/registros",
        data={"meta_id": str(meta_horas), "descripcion": "Avance uno", "horas": "4"},
    )
    panel = client.get("/")
    assert "Mis metas y su avance" in panel.text
    assert "Cierre contable octubre" in panel.text
    assert "40%" in panel.text

    # Sin horas estimadas: 1 de 2 reportes completados => 50%.
    client.post(
        "/registros",
        data={
            "meta_id": str(meta_sin),
            "descripcion": "Paso uno",
            "estado": "completado",
        },
    )
    client.post(
        "/registros",
        data={
            "meta_id": str(meta_sin),
            "descripcion": "Paso dos",
            "fecha": (date.today() - timedelta(days=1)).isoformat(),
            "estado": "en_progreso",
        },
    )
    assert "50%" in client.get("/").text

    # La analista no entra a la vista del jefe.
    assert client.get(
        "/metas/progreso", follow_redirects=False
    ).headers["location"].startswith("/?msg=")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303

    pagina = client.get("/metas/progreso")
    assert pagina.status_code == 200
    assert "Progreso de metas" in pagina.text
    assert "40%" in pagina.text
    assert "50%" in pagina.text
    assert "10 h" in pagina.text  # horas estimadas de la meta

    # La lista de metas tambien muestra el avance.
    assert "40%" in client.get("/metas").text

    # 7 horas mas => 11 de 10 previstas, el avance se topa en 100%.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    client.post(
        "/registros",
        data={"meta_id": str(meta_horas), "descripcion": "Avance dos", "horas": "7"},
    )
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    assert "100%" in client.get("/metas/progreso").text

    # Cerrar una meta la deja en 100%.
    assert client.post(
        f"/metas/{meta_sin}/estado", data={"estado": "cerrada"}, follow_redirects=False
    ).status_code == 303
    progreso = client.get("/metas/progreso?ver=todas")
    assert progreso.status_code == 200
    assert "Meta sin estimacion" in progreso.text

    # La vista "activas" deja afuera la meta cerrada.
    assert "Meta sin estimacion" not in client.get("/metas/progreso?ver=activas").text


def test_migracion_agrega_columnas_a_una_base_vieja():
    import sqlite3

    from sqlalchemy import create_engine

    carpeta = Path(tempfile.mkdtemp(prefix="cg-migracion-"))
    archivo = carpeta / "viejo.db"

    conexion = sqlite3.connect(archivo)
    conexion.execute("CREATE TABLE registros (id INTEGER PRIMARY KEY, descripcion TEXT)")
    conexion.execute("CREATE TABLE metas (id INTEGER PRIMARY KEY, titulo TEXT)")
    conexion.commit()
    conexion.close()

    motor = create_engine(f"sqlite:///{archivo}")
    original = migraciones.engine
    migraciones.engine = motor
    try:
        migraciones.aplicar_migraciones()
    finally:
        migraciones.engine = original

    conexion = sqlite3.connect(archivo)
    try:
        columnas = {fila[1] for fila in conexion.execute("PRAGMA table_info(registros)")}
        columnas_metas = {
            fila[1] for fila in conexion.execute("PRAGMA table_info(metas)")
        }
    finally:
        conexion.close()
    assert {"comentario", "comentado_en"} <= columnas
    assert "horas_estimadas" in columnas_metas


def test_historial_por_persona(client):
    assert login(client, "jefe", "1234").status_code == 303
    maria_id = _id_usuario("mperez")

    pagina = client.get(f"/personas/{maria_id}")
    assert pagina.status_code == 200
    assert "Maria Perez" in pagina.text
    assert "Historial de reportes" in pagina.text
    assert "Revisadas 40 facturas pendientes" in pagina.text
    assert "Metas de Maria" in pagina.text

    # El filtro por fechas tambien funciona aca.
    vacio = client.get(
        f"/personas/{maria_id}", params={"desde": "2000-01-01", "hasta": "2000-01-31"}
    )
    assert "No hay reportes con esos filtros" in vacio.text

    # Usuario inexistente.
    assert client.get(
        "/personas/9999", follow_redirects=False
    ).headers["location"].startswith("/equipo?msg=")

    # Un empleado no entra al historial de nadie.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert client.get(
        f"/personas/{maria_id}", follow_redirects=False
    ).headers["location"].startswith("/?msg=")


def test_grafico_de_horas_en_el_panel(client):
    assert login(client, "jefe", "1234").status_code == 303
    panel = client.get("/")
    assert "Horas por semana" in panel.text
    assert 'class="col c1"' in panel.text
    assert "La barra más alta" in panel.text

    # El empleado no ve el grafico del equipo.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert "Horas por semana" not in client.get("/").text


def test_horas_por_semana_escala_las_barras():
    from app.metricas import horas_por_semana

    db = SessionLocal()
    try:
        equipo = db.query(Usuario).order_by(Usuario.nombre).all()
        datos = horas_por_semana(db, equipo, date.today(), semanas=3)
    finally:
        db.close()

    assert len(datos["bloques"]) == 3
    assert datos["bloques"][-1]["actual"] is True
    assert datos["bloques"][0]["actual"] is False

    altos = [barra["alto"] for bloque in datos["bloques"] for barra in bloque["barras"]]
    assert max(altos) == 100  # la barra mas alta siempre llega al tope
    assert all(0 <= alto <= 100 for alto in altos)
    assert datos["bloques"][-1]["total"] > 0  # la semana en curso tiene horas
    assert datos["hay_datos"] is True


def test_ips_locales_incluye_el_hotspot(monkeypatch):
    from app import security

    class Resultado:
        stdout = (
            "lo: flags=73<UP,LOOPBACK,RUNNING>\n"
            "        inet 127.0.0.1  netmask 255.0.0.0\n"
            "wlan0: flags=4163<UP,BROADCAST,RUNNING,MULTICAST>\n"
            "        inet 192.168.1.47  netmask 255.255.255.0\n"
            "ap0: flags=4163<UP,BROADCAST,RUNNING,MULTICAST>\n"
            "        inet 10.187.127.38  netmask 255.255.255.0\n"
        )

    monkeypatch.setattr(security.subprocess, "run", lambda *a, **k: Resultado())
    ips = security.ips_locales()

    assert "192.168.1.47" in ips  # la Wi-Fi a la que esta conectado
    assert "10.187.127.38" in ips  # el hotspot que comparte
    assert "127.0.0.1" not in ips  # el loopback no sirve para la red


def test_ips_locales_acepta_el_formato_viejo_de_ifconfig(monkeypatch):
    from app import security

    class Resultado:
        stdout = (
            "eth0: flags=4163<UP,BROADCAST,RUNNING>\n"
            "        inet addr:192.168.0.10  Bcast:192.168.0.255  Mask:255.255.255.0\n"
        )

    monkeypatch.setattr(security.subprocess, "run", lambda *a, **k: Resultado())
    assert "192.168.0.10" in security.ips_locales()


def test_ips_locales_sin_ifconfig_no_falla(monkeypatch):
    from app import security

    def explota(*args, **kwargs):
        raise FileNotFoundError("ifconfig no existe")

    monkeypatch.setattr(security.subprocess, "run", explota)
    assert isinstance(security.ips_locales(), list)


def test_guia_para_el_equipo(client):
    assert login(client, "mperez", "4821").status_code == 303
    guia = client.get("/guia")
    assert guia.status_code == 200
    assert "Registrar lo que hiciste" in guia.text
    assert "en progreso" in guia.text
    assert "Cambiar mi PIN" in guia.text


def test_reporte_pdf_incluye_el_avance(client, monkeypatch):
    import fpdf

    original = fpdf.FPDF.__init__

    def sin_compresion(self, *args, **kwargs):
        original(self, *args, **kwargs)
        self.set_compression(False)  # deja el texto legible para poder verificarlo

    monkeypatch.setattr(fpdf.FPDF, "__init__", sin_compresion)

    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={"titulo": "Meta del PDF", "horas_estimadas": "8"},
        follow_redirects=False,
    )

    respuesta = client.get("/reportes/mensual.pdf")
    assert respuesta.status_code == 200

    texto = respuesta.content.decode("latin-1")
    assert "Avance promedio de metas activas" in texto
    assert "Metas con plazo vencido" in texto
    assert "Meta del PDF" in texto
    # Encabezados de la tabla de metas, incluida la columna nueva de avance.
    for encabezado in ("Asignada a", "Plazo", "Avances mes", "Estado"):
        assert encabezado in texto


def test_exportar_excel_con_filtros(client):
    assert login(client, "jefe", "1234").status_code == 303

    respuesta = client.get("/registros/exportar.xlsx")
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in respuesta.headers["content-disposition"]
    assert respuesta.content[:2] == b"PK"  # un .xlsx es un zip

    hoja = load_workbook(BytesIO(respuesta.content)).active
    assert hoja["A1"].value == "Fecha"
    textos = [c.value for fila in hoja.iter_rows() for c in fila if c.value]
    assert "Revisadas 40 facturas pendientes" in textos
    assert "TOTAL HORAS" in textos

    # Filtrando por una persona sin registros, el archivo sale sin datos.
    filtrado = client.get("/registros/exportar.xlsx", params={"usuario_id": "999"})
    valores = [
        c.value
        for fila in load_workbook(BytesIO(filtrado.content)).active.iter_rows()
        for c in fila
        if c.value
    ]
    assert "Revisadas 40 facturas pendientes" not in valores

    # Un empleado no puede exportar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    assert client.get(
        "/registros/exportar.xlsx", follow_redirects=False
    ).headers["location"].startswith("/?msg=")


def test_redireccion_sin_json_en_el_cuerpo(client):
    """La redireccion de seguridad responde 303 con Location y sin cuerpo."""
    respuesta = client.get("/", follow_redirects=False)
    assert respuesta.status_code == 303
    assert respuesta.headers["location"] == "/login"
    assert "application/json" not in respuesta.headers.get("content-type", "")
    assert respuesta.content == b""


def test_clave_de_sesion_propia_por_instalacion(tmp_path, monkeypatch):
    """Sin CG_SECRET_KEY se genera una clave por instalacion y se reutiliza."""
    monkeypatch.delenv("CG_SECRET_KEY", raising=False)
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)

    primera = config._clave_de_sesion()
    assert len(primera) >= 32
    assert config._clave_de_sesion() == primera
    assert (tmp_path / ".secret_key").read_text(encoding="utf-8") == primera

    # Si el ambiente define la clave, manda la del ambiente.
    monkeypatch.setenv("CG_SECRET_KEY", "la-que-diga-el-ambiente")
    assert config._clave_de_sesion() == "la-que-diga-el-ambiente"


def test_aviso_si_el_respaldo_queda_en_el_mismo_disco(client):
    # Por defecto los respaldos van a data/respaldos: mismo disco que la base.
    assert backup.respaldos_en_mismo_disco() is True
    assert login(client, "jefe", "1234").status_code == 303
    assert "mismo disco" in client.get("/respaldos").text


def test_restaurar_respaldo(client):
    assert login(client, "jefe", "1234").status_code == 303

    client.post(
        "/metas", data={"titulo": "Meta que va a volver"}, follow_redirects=False
    )
    client.post("/respaldos/ahora", follow_redirects=False)
    respaldo = max(
        config.BACKUP_DIR.glob("control_*_manual.db"),
        key=lambda ruta: ruta.stat().st_mtime,
    )

    meta_id = _id_meta("Meta que va a volver")
    client.post(f"/metas/{meta_id}/eliminar", follow_redirects=False)

    db = SessionLocal()
    try:
        archivada = db.query(Meta).filter_by(titulo="Meta que va a volver").one()
        assert archivada.estado == "eliminada"
    finally:
        db.close()

    respuesta = client.post(
        "/respaldos/restaurar", data={"nombre": respaldo.name}, follow_redirects=False
    )
    assert respuesta.status_code == 303
    assert "Base restaurada" in unquote(respuesta.headers["location"])

    db = SessionLocal()
    try:
        assert db.query(Meta).filter_by(titulo="Meta que va a volver").count() == 1
    finally:
        db.close()

    # La pagina lista el respaldo y la copia de seguridad previa.
    assert "antes-de-restaurar" in client.get("/respaldos").text

    # Un empleado no puede restaurar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    denegado = client.post(
        "/respaldos/restaurar", data={"nombre": respaldo.name}, follow_redirects=False
    )
    assert denegado.status_code == 303
    assert denegado.headers["location"].startswith("/?msg=")


def test_restaurar_rechaza_archivos_que_no_son_respaldos(client):
    assert login(client, "jefe", "1234").status_code == 303

    (config.BACKUP_DIR / "pepito.txt").write_text("hola", encoding="utf-8")
    (config.BACKUP_DIR / "control_basura.db").write_text(
        "esto no es sqlite", encoding="utf-8"
    )

    casos = [
        "../../control.db",  # fuera de la carpeta de respaldos
        "no_existe.db",  # no esta
        "pepito.txt",  # esta pero no tiene formato de respaldo
        "control_basura.db",  # tiene el formato pero esta corrupto
    ]
    for nombre in casos:
        respuesta = client.post(
            "/respaldos/restaurar", data={"nombre": nombre}, follow_redirects=False
        )
        assert respuesta.status_code == 303
        assert "No se pudo restaurar" in unquote(respuesta.headers["location"]), nombre


def _ultimo_registro_de(login_usuario: str) -> int:
    db = SessionLocal()
    try:
        persona = db.query(Usuario).filter_by(usuario=login_usuario).one()
        registro = (
            db.query(Registro)
            .filter(Registro.usuario_id == persona.id)
            .order_by(Registro.id.desc())
            .first()
        )
        return registro.id
    finally:
        db.close()


def test_filtrar_registros_por_meta(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas", data={"titulo": "Meta del filtro A", "horas_estimadas": "10"}
    )
    client.post(
        "/metas", data={"titulo": "Meta del filtro B", "horas_estimadas": "10"}
    )
    meta_a = _id_meta("Meta del filtro A")
    meta_b = _id_meta("Meta del filtro B")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    for meta_id, texto in ((meta_a, "Carga de la meta A"), (meta_b, "Carga de la meta B")):
        assert (
            client.post(
                "/registros",
                data={"meta_id": str(meta_id), "descripcion": texto, "horas": "2"},
                follow_redirects=False,
            ).status_code
            == 303
        )

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303

    todo = client.get("/registros").text
    assert "Carga de la meta A" in todo
    assert "Carga de la meta B" in todo

    # Filtrando por meta A solo queda esa carga y el combo marca la seleccion.
    pagina = client.get("/registros", params={"meta_id": str(meta_a)})
    assert pagina.status_code == 200
    assert "Carga de la meta A" in pagina.text
    assert "Carga de la meta B" not in pagina.text
    assert f'value="{meta_a}" selected' in pagina.text

    # La exportacion respeta el mismo filtro.
    hoja = load_workbook(
        BytesIO(
            client.get(
                "/registros/exportar.xlsx", params={"meta_id": str(meta_a)}
            ).content
        )
    ).active
    textos = [c.value for fila in hoja.iter_rows() for c in fila if c.value]
    assert "Carga de la meta A" in textos
    assert "Carga de la meta B" not in textos

    # Un meta_id que no es numero se ignora en vez de romper.
    assert client.get("/registros", params={"meta_id": "abc"}).status_code == 200


def test_duplicar_meta(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={
            "titulo": "Corte de septiembre 2026",
            "descripcion": "Cierra las cuentas del mes",
            "fecha_inicio": "2026-09-01",
            "fecha_limite": "2026-09-30",
            "horas_estimadas": "12",
        },
    )
    original = _id_meta("Corte de septiembre 2026")

    # El original se cierra: la copia tiene que arrancar activa igual.
    client.post(f"/metas/{original}/estado", data={"estado": "cerrada"})

    respuesta = client.post(f"/metas/{original}/duplicar", follow_redirects=False)
    assert respuesta.status_code == 303
    destino = respuesta.headers["location"]
    assert destino.startswith("/metas/") and "/editar?msg=" in destino
    assert "Meta%20duplicada" in destino

    db = SessionLocal()
    try:
        copia = db.get(Meta, int(destino.split("/")[2]))
        assert copia.id != original
        assert copia.titulo == "Corte de septiembre 2026 (copia)"
        assert copia.descripcion == "Cierra las cuentas del mes"
        assert copia.fecha_inicio == date(2026, 9, 1)
        assert copia.fecha_limite == date(2026, 9, 30)
        assert copia.horas_estimadas == 12
        assert copia.estado == "activa"
        assert db.get(Meta, original).estado == "cerrada"
    finally:
        db.close()

    assert "Corte de septiembre 2026 (copia)" in client.get("/metas").text

    # Un empleado no puede duplicar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    denegado = client.post(f"/metas/{original}/duplicar", follow_redirects=False)
    assert denegado.status_code == 303
    assert denegado.headers["location"].startswith("/?msg=")

    # Archivar no borra: la meta queda en pie y de paso se puede duplicar.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    borrada = client.post(f"/metas/{original}/eliminar", follow_redirects=False)
    assert borrada.status_code == 303

    db = SessionLocal()
    try:
        assert db.get(Meta, original).estado == "eliminada"
    finally:
        db.close()

    copia_archivada = client.post(f"/metas/{original}/duplicar", follow_redirects=False)
    assert copia_archivada.status_code == 303
    assert "Meta%20duplicada" in copia_archivada.headers["location"]

    # Duplicar algo que ya no existe no tira error.
    inexistente = client.post("/metas/999999/duplicar", follow_redirects=False)
    assert inexistente.status_code == 303
    assert "Esa%20meta%20ya%20no%20existe" in inexistente.headers["location"]


def test_repetir_carga_del_dia_anterior(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post("/metas", data={"titulo": "Meta para repetir", "horas_estimadas": "8"})
    meta_id = _id_meta("Meta para repetir")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    client.post(
        "/registros",
        data={
            "meta_id": str(meta_id),
            "descripcion": "Trabajo de ayer para repetir",
            "fecha": (date.today() - timedelta(days=1)).isoformat(),
            "horas": "6.5",
            "estado": "completado",
        },
    )
    registro_id = _ultimo_registro_de("mperez")

    # El formulario queda precargado con lo anterior pero la fecha es hoy.
    panel = client.get("/", params={"repetir": str(registro_id)})
    assert panel.status_code == 200
    assert "Repitiendo el registro del" in panel.text
    assert "required>Trabajo de ayer para repetir</textarea>" in panel.text
    assert 'value="6.5"' in panel.text
    assert f'value="{meta_id}" selected' in panel.text
    assert f'value="{date.today().isoformat()}"' in panel.text
    assert f'href="/?repetir={registro_id}"' in panel.text

    # Se guarda como un registro nuevo del dia de hoy.
    assert (
        client.post(
            "/registros",
            data={
                "meta_id": str(meta_id),
                "descripcion": "Trabajo de ayer para repetir",
                "horas": "6.5",
                "estado": "completado",
            },
            follow_redirects=False,
        ).headers["location"]
        == "/?msg=Registro%20guardado"
    )
    db = SessionLocal()
    try:
        persona = db.query(Usuario).filter_by(usuario="mperez").one()
        hoyos = (
            db.query(Registro)
            .filter(
                Registro.usuario_id == persona.id,
                Registro.fecha == date.today(),
                Registro.meta_id == meta_id,
            )
            .count()
        )
        totales = (
            db.query(Registro)
            .filter(
                Registro.usuario_id == persona.id,
                Registro.meta_id == meta_id,
            )
            .count()
        )
        # Solo la copia es de hoy: la original quedo en el dia anterior.
        assert hoyos == 1
        assert totales == 2
    finally:
        db.close()

    # Un id inexistente o ajeno no precarga nada ni rompe la pagina.
    assert "Repitiendo el registro del" not in client.get(
        "/", params={"repetir": "999999"}
    ).text
    assert "Repitiendo el registro del" not in client.get(
        "/", params={"repetir": "no-numero"}
    ).text

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    ajeno = client.get("/", params={"repetir": str(registro_id)})
    assert ajeno.status_code == 200
    assert "Repitiendo el registro del" not in ajeno.text
    # El jefe solo ve el boton de repetir en sus propios registros.
    assert f'href="/?repetir={registro_id}"' not in ajeno.text


def test_calendario_de_carga(client):
    hoy = date.today()

    db = SessionLocal()
    try:
        equipo = (
            db.query(Usuario).filter_by(activo=True).order_by(Usuario.nombre).all()
        )
        cal = calendario_mes(db, equipo, hoy.year, hoy.month)
        assert cal["huecos"] == date(hoy.year, hoy.month, 1).weekday()
        assert len(cal["dias"]) == calendar.monthrange(hoy.year, hoy.month)[1]
        assert cal["total_horas"] >= 0
        del_hoy = next(d for d in cal["dias"] if d["fecha"] == hoy)
        assert len(del_hoy["celdas"]) == len(equipo)
        assert all("iniciales" in celda for celda in del_hoy["celdas"])
    finally:
        db.close()

    # Un empleado no entra.
    assert login(client, "mperez", "4821").status_code == 303
    denegado = client.get("/calendario", follow_redirects=False)
    assert denegado.status_code == 303
    assert denegado.headers["location"].startswith("/?msg=")
    client.post("/logout", follow_redirects=False)

    assert login(client, "jefe", "1234").status_code == 303
    pagina = client.get("/calendario", params={"anio": hoy.year, "mes": hoy.month})
    assert pagina.status_code == 200
    assert "Calendario de carga" in pagina.text
    assert ">Lun<" in pagina.text
    assert ">Dom<" in pagina.text
    assert "Mes anterior" in pagina.text
    assert "Mes siguiente" in pagina.text
    assert "cal-body" in pagina.text

    # Se puede ir a cualquier mes.
    agosto = client.get("/calendario", params={"anio": 2026, "mes": 8})
    assert agosto.status_code == 200
    assert '<option value="8" selected>agosto</option>' in agosto.text
    db = SessionLocal()
    try:
        equipo = (
            db.query(Usuario).filter_by(activo=True).order_by(Usuario.nombre).all()
        )
        agosto_cal = calendario_mes(db, equipo, 2026, 8)
    finally:
        db.close()
    # Celdas dibujadas = huecos del lunes + dias del mes.
    assert agosto.text.count('class="cal-dia') == (
        agosto_cal["huecos"] + 31
    )

    # Periodo raro: si no es numero vuelve al mes en curso, si esta fuera de
    # rango avisa.
    assert client.get(
        "/calendario", params={"anio": "abc"}, follow_redirects=False
    ).status_code == 200
    fuera = client.get(
        "/calendario", params={"mes": "13"}, follow_redirects=False
    )
    assert fuera.status_code == 303
    assert fuera.headers["location"].startswith("/calendario?msg=")


# --------------------------------------------------------------------------
# Fase 1: base de datos (WAL, indices, totales, respaldos)
# --------------------------------------------------------------------------
def test_base_trabaja_en_modo_wal():
    from sqlalchemy import text

    db = SessionLocal()
    try:
        modo = db.execute(text("PRAGMA journal_mode")).scalar()
    finally:
        db.close()
    assert modo == "wal"


def test_indices_creados_en_la_base():
    import sqlite3

    conexion = sqlite3.connect(str(config.DATA_DIR / "control.db"))
    try:
        nombres = {
            fila[0]
            for fila in conexion.execute(
                "SELECT name FROM sqlite_master WHERE type = 'index'"
            )
        }
    finally:
        conexion.close()
    assert {
        "ix_registros_fecha",
        "ix_registros_usuario_fecha",
        "ix_registros_unico",
        "ix_metas_estado_fecha_limite",
    } <= nombres


def test_totales_no_se_cortan_con_el_limite_de_la_lista(client):
    """La lista trae 500 filas, pero el total tiene que contar todas."""
    db = SessionLocal()
    creados = []
    try:
        jefe = db.query(Usuario).filter_by(usuario="jefe").one()
        existentes = db.query(Registro).all()
        antes = len(existentes)
        horas_antes = sum(r.horas for r in existentes)
        for i in range(505):
            registro = Registro(
                usuario_id=jefe.id,
                fecha=date.today() - timedelta(days=i % 30),
                descripcion=f"trabajo de prueba {i}",
                horas=1.0,
                estado="completado",
            )
            db.add(registro)
            creados.append(registro)
        db.commit()
        ids = [r.id for r in creados]
    finally:
        db.close()

    try:
        assert login(client, "jefe", "1234").status_code == 303
        pagina = client.get("/registros")
        assert pagina.status_code == 200
        assert f"{antes + 505} registro(s)" in pagina.text
        assert f"{horas_antes + 505.0:.1f} horas acumuladas" in pagina.text
        assert "La vista trae los últimos 500" in pagina.text
    finally:
        db = SessionLocal()
        try:
            db.query(Registro).filter(Registro.id.in_(ids)).delete(
                synchronize_session=False
            )
            db.commit()
        finally:
            db.close()


def test_exportar_no_hace_una_consulta_por_fila(client):
    """La exportacion no trae metas ni usuarios aparte: sin eager loading
    cada fila dispararia su propia consulta para leer la meta."""
    from sqlalchemy import event

    from app.database import engine

    db = SessionLocal()
    creados = []
    try:
        jefe = db.query(Usuario).filter_by(usuario="jefe").one()
        for i in range(40):
            meta = Meta(
                titulo=f"Meta carga rapida {i}",
                fecha_limite=date.today(),
                horas_estimadas=0.0,
            )
            db.add(meta)
            db.flush()
            registro = Registro(
                usuario_id=jefe.id,
                meta_id=meta.id,
                fecha=date.today(),
                descripcion=f"carga rapida {i}",
                horas=1.0,
                estado="en_progreso",
            )
            db.add(registro)
            creados.append(registro)
        db.commit()
        ids_registros = [r.id for r in creados]
        ids_metas = [r.meta_id for r in creados]
    finally:
        db.close()

    consultas: list[str] = []

    def contar(_conexion, _cursor, sentencia, _parametros, _contexto, _muchos):
        consultas.append(sentencia)

    assert login(client, "jefe", "1234").status_code == 303
    event.listen(engine, "before_cursor_execute", contar)
    try:
        export = client.get("/registros/exportar.xlsx")
    finally:
        event.remove(engine, "before_cursor_execute", contar)
        db = SessionLocal()
        try:
            db.query(Registro).filter(Registro.id.in_(ids_registros)).delete(
                synchronize_session=False
            )
            db.query(Meta).filter(Meta.id.in_(ids_metas)).delete(
                synchronize_session=False
            )
            db.commit()
        finally:
            db.close()

    assert export.status_code == 200
    assert len(consultas) < 15


def test_respaldo_invalido_se_borra_y_no_queda_en_la_lista(client, monkeypatch):
    carpeta = backup.asegurar_carpeta()

    def romper(_ruta):
        raise ValueError("el archivo esta corrupto")

    monkeypatch.setattr(backup, "_verificar_sqlite", romper)
    with pytest.raises(ValueError):
        backup.crear_respaldo(etiqueta="malo")

    assert not list(carpeta.glob("*_malo.db"))


def test_borrar_archivos_wal_quita_los_dos(tmp_path, monkeypatch):
    base = tmp_path / "control.db"
    monkeypatch.setattr(backup, "DB_PATH", base)
    wal = tmp_path / "control.db-wal"
    shm = tmp_path / "control.db-shm"
    wal.write_text("viejo", encoding="utf-8")
    shm.write_text("viejo", encoding="utf-8")

    backup._borrar_archivos_wal()

    assert not wal.exists()
    assert not shm.exists()


def test_restaurar_saca_los_archivos_wal(client, monkeypatch):
    assert login(client, "jefe", "1234").status_code == 303
    client.post("/respaldos/ahora", follow_redirects=False)
    respaldo = max(
        config.BACKUP_DIR.glob("control_*_manual.db"),
        key=lambda ruta: ruta.stat().st_mtime,
    )

    llamadas = []
    monkeypatch.setattr(
        backup, "_borrar_archivos_wal", lambda: llamadas.append(True)
    )
    respuesta = client.post(
        "/respaldos/restaurar", data={"nombre": respaldo.name}, follow_redirects=False
    )

    assert respuesta.status_code == 303
    assert llamadas


# --------------------------------------------------------------------------
# Integridad de datos (Fase 2)
# --------------------------------------------------------------------------
def _usuario_id(usuario: str) -> int:
    db = SessionLocal()
    try:
        return db.query(Usuario).filter_by(usuario=usuario).one().id
    finally:
        db.close()


def test_no_se_puede_cargar_en_el_futuro(client):
    assert login(client, "jefe", "1234").status_code == 303
    respuesta = client.post(
        "/registros",
        data={
            "descripcion": "Tarea de manana",
            "fecha": (date.today() + timedelta(days=1)).isoformat(),
        },
        follow_redirects=False,
    )
    assert "fecha%20futura" in respuesta.headers["location"]


def test_no_se_puede_cargar_fuera_del_rango(client):
    assert login(client, "jefe", "1234").status_code == 303
    tarde = (date.today() - timedelta(days=config.DIAS_ATRASO + 1)).isoformat()
    respuesta = client.post(
        "/registros",
        data={"descripcion": "Del ano pasado", "fecha": tarde},
        follow_redirects=False,
    )
    assert "Fuera%20de%20rango" in respuesta.headers["location"]

    # Justo en el limite: se acepta.
    limite = (date.today() - timedelta(days=config.DIAS_ATRASO)).isoformat()
    ok = client.post(
        "/registros",
        data={"descripcion": "Ultimo dia valido", "fecha": limite},
        follow_redirects=False,
    )
    assert "Registro%20guardado" in ok.headers["location"]


def test_un_solo_reporte_por_dia_y_meta(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post("/metas", data={"titulo": "Meta diaria"})
    meta_id = _id_meta("Meta diaria")

    primera = client.post(
        "/registros",
        data={"meta_id": str(meta_id), "descripcion": "un trabajo"},
        follow_redirects=False,
    )
    assert "Registro%20guardado" in primera.headers["location"]

    duplicado = client.post(
        "/registros",
        data={"meta_id": str(meta_id), "descripcion": "otro trabajo"},
        follow_redirects=False,
    )
    assert "Ya%20hay%20un%20registro" in duplicado.headers["location"]

    # Sin meta tambien aplica, aunque el indice unico no cubre los NULL.
    dia_prueba = date.today() - timedelta(days=25)
    db = SessionLocal()
    try:
        jefe = _usuario_id("jefe")
        db.query(Registro).filter_by(
            usuario_id=jefe, fecha=dia_prueba, meta_id=None
        ).delete()
        db.commit()
    finally:
        db.close()

    libre = client.post(
        "/registros",
        data={"descripcion": "sin meta", "fecha": dia_prueba.isoformat()},
        follow_redirects=False,
    )
    assert "Registro%20guardado" in libre.headers["location"]
    libre_dup = client.post(
        "/registros",
        data={"descripcion": "sin meta otra vez", "fecha": dia_prueba.isoformat()},
        follow_redirects=False,
    )
    assert "Ya%20hay%20un%20registro" in libre_dup.headers["location"]


def test_no_se_puede_reportar_sobre_una_meta_ajena(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post(
        "/metas",
        data={"titulo": "Meta solo del jefe", "asignado_a": str(_usuario_id("jefe"))},
    )
    client.post(
        "/metas",
        data={"titulo": "Meta solo de Maria", "asignado_a": str(_usuario_id("mperez"))},
    )
    ajena = _id_meta("Meta solo del jefe")
    propia = _id_meta("Meta solo de Maria")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303

    denegado = client.post(
        "/registros",
        data={"meta_id": str(ajena), "descripcion": "me meto en la ajena"},
        follow_redirects=False,
    )
    assert "no%20te%20corresponde" in denegado.headers["location"]

    propia_ok = client.post(
        "/registros",
        data={"meta_id": str(propia), "descripcion": "mi trabajo"},
        follow_redirects=False,
    )
    assert "Registro%20guardado" in propia_ok.headers["location"]

    # Una meta sin asignar es de todo el equipo.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    client.post("/metas", data={"titulo": "Meta de todos"})
    equipo = _id_meta("Meta de todos")
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    libre = client.post(
        "/registros",
        data={"meta_id": str(equipo), "descripcion": "para todos"},
        follow_redirects=False,
    )
    assert "Registro%20guardado" in libre.headers["location"]


def test_archivar_meta_conserva_los_reportes(client):
    assert login(client, "jefe", "1234").status_code == 303
    client.post("/metas", data={"titulo": "Meta a archivar"})
    meta_id = _id_meta("Meta a archivar")

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    client.post(
        "/registros",
        data={"meta_id": str(meta_id), "descripcion": "avance"},
        follow_redirects=False,
    )

    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    archivar = client.post(f"/metas/{meta_id}/eliminar", follow_redirects=False)
    assert archivar.status_code == 303
    assert "archivada" in unquote(archivar.headers["location"])

    db = SessionLocal()
    try:
        meta = db.get(Meta, meta_id)
        assert meta.estado == "eliminada"
        assert db.query(Registro).filter_by(meta_id=meta_id).count() == 1
    finally:
        db.close()

    # Sigue visible para el jefe, marcada y con la opcion de restaurar.
    pagina = client.get("/metas")
    assert "archivada" in pagina.text
    assert "Restaurar" in pagina.text

    # El empleado ya no la puede elegir para cargar (en su lista de metas si
    # aparece, porque el reporte viejo conserva el vinculo).
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "mperez", "4821").status_code == 303
    panel = client.get("/")
    assert f'<option value="{meta_id}"' not in panel.text
    assert "Meta a archivar" in panel.text

    # Se restaura y vuelve a aparecer.
    assert client.post("/logout", follow_redirects=False).status_code == 303
    assert login(client, "jefe", "1234").status_code == 303
    restaurar = client.post(
        f"/metas/{meta_id}/estado", data={"estado": "activa"}, follow_redirects=False
    )
    assert restaurar.status_code == 303

    db = SessionLocal()
    try:
        assert db.get(Meta, meta_id).estado == "activa"
    finally:
        db.close()
    assert "Meta a archivar" in client.get("/metas").text


def test_el_indice_unico_no_rompe_si_hay_duplicados_viejos():
    """Una base de la oficina puede traer duplicados previos: avisa y sigue."""
    from sqlalchemy import create_engine

    from app.migraciones import _asegurar_indice_unico

    motor = create_engine("sqlite://")
    with motor.begin() as conexion:
        conexion.exec_driver_sql("CREATE TABLE usuarios (id INTEGER, nombre TEXT)")
        conexion.exec_driver_sql("INSERT INTO usuarios VALUES (1, 'Maria')")
        conexion.exec_driver_sql(
            "CREATE TABLE registros (usuario_id INTEGER, fecha DATE, meta_id INTEGER)"
        )
        conexion.exec_driver_sql("INSERT INTO registros VALUES (1, '2026-09-01', 7)")
        conexion.exec_driver_sql("INSERT INTO registros VALUES (1, '2026-09-01', 7)")
        _asegurar_indice_unico(conexion)
        nombres = {
            fila[0]
            for fila in conexion.exec_driver_sql(
                "SELECT name FROM sqlite_master WHERE type='index'"
            )
        }
        assert "ix_registros_unico" not in nombres

    motor = create_engine("sqlite://")
    with motor.begin() as conexion:
        conexion.exec_driver_sql("CREATE TABLE usuarios (id INTEGER, nombre TEXT)")
        conexion.exec_driver_sql(
            "CREATE TABLE registros (usuario_id INTEGER, fecha DATE, meta_id INTEGER)"
        )
        _asegurar_indice_unico(conexion)
        nombres = {
            fila[0]
            for fila in conexion.exec_driver_sql(
                "SELECT name FROM sqlite_master WHERE type='index'"
            )
        }
        assert "ix_registros_unico" in nombres
